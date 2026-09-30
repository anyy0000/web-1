"""결과 저장: Google 스프레드시트 / 엑셀.

매 실행마다 기존 행을 읽어 신규 행과 합친 뒤, 구분 → 마감일(D-day) 순으로 정렬해 다시 쓴다.
상태/메모처럼 사람이 적은 칸과 모르는 열도 그대로 보존한다.
  - 메인 탭 "링커리어": 매칭된 공고
  - "_seen" 탭: 이미 확인한 공고 ID (제외된 것 포함 → 다음 실행 때 다시 읽지 않음)
  - "설정" 탭 (구글시트만): 관심기업/관심분야/키워드를 시트에서 직접 추가
"""

from __future__ import annotations

import json
import os
from datetime import date, datetime, timedelta
from pathlib import Path

from crawler import KST

HEADER = ["구분", "D-day", "추천도", "제목", "주최", "기관유형", "분야", "마감일",
          "내용요약", "준비할 것", "참가대상", "추천이유", "시상(만원)", "링크", "상태", "메모", "수집일", "ID"]
RENAMED = {"매칭근거": "추천이유"}  # 이전 버전 열 이름
DROPPED = {"점수"}
DATE_COLS = {"마감일", "수집일"}
DEADLINE_COL = chr(ord("A") + HEADER.index("마감일"))
MAIN_TAB = "링커리어"
SEEN_TAB = "_seen"
SETTINGS_TAB = "설정"
SETTINGS_HEADER = ["구분", "값", "메모"]
# 설정 탭 구분 ↔ config.yaml 키 (처음 한 번 config.yaml 내용을 시트로 옮긴다)
SETTING_KINDS = {
    "관심기업": "organizations", "관심분야": "categories", "참고분야": "sub_categories",
    "주최키워드": "organizer_keywords", "키워드": "content_keywords",
    "제외키워드": "exclude_keywords", "제외기업": None,
}
GUIDE_MARK = "조건 설명 (v2)"
GUIDE = [
    [GUIDE_MARK, "점수", "설명"],
    ["관심기업", "+4 → 추천", "주최사 이름에 이 글자가 들어가면. 일부만 적어도 됨 (예: 'CJ' → CJ제일제당·CJ프레시웨이 모두)"],
    ["관심분야", "+4 → 추천", "링커리어 분야와 정확히 같을 때. 메인 탭 '분야' 열에 나온 이름을 복사해서 적기"],
    ["참고분야", "+2 → 검토", "링커리어 분야와 정확히 같을 때. 다른 조건과 합쳐 4점 이상이면 추천"],
    ["주최키워드", "+3", "주최사 이름에 들어가면 (업종 추정용: 식품, 제약, 헬스 …)"],
    ["키워드", "제목 +2 / 본문 +1", "제목은 최대 +4, 본문(포스터 OCR 포함)은 최대 +3"],
    ["제외키워드", "-3", "제목·본문에 들어가면 감점 (예: '금융 건강')"],
    ["제외기업", "항상 제외", "이 주최사 공고는 점수와 관계없이 제외"],
    ["", "", ""],
    ["판정", "", "합계 4점 이상 = 추천, 2~3점 = 검토, 그 미만은 시트에 넣지 않음"],
    ["사용법", "", "A~B열에 한 줄씩 '구분'과 '값'을 적거나 지우면 다음 실행부터 반영. 빈 줄은 무시"],
    ["적용 범위", "", "이후 새로 올라오는 공고부터 적용. 이미 제외된 공고도 다시 보려면 _seen 탭 2행부터 아래를 지우기"],
    ["분야 예시", "", "요리/식품, 의료/보건, 체육/헬스, 뷰티/미용/화장품, 기획/아이디어, 광고/마케팅, 서포터즈, 봉사활동, 과학/공학"],
    ["re:로 시작", "", "특수 패턴(정규식)입니다. 예: '대상'이 '참가대상'과 헷갈리지 않게 막는 용도라 그대로 두세요"],
]
# 행 위치와 무관하게 동작 (정렬해도 깨지지 않음)
DDAY_FORMULA = f'=IF(INDIRECT("{DEADLINE_COL}"&ROW())="","",DAYS(INDIRECT("{DEADLINE_COL}"&ROW()),TODAY()))'


def to_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 20000 < value < 80000:  # 스프레드시트 날짜 일련번호
        return date(1899, 12, 30) + timedelta(days=int(value))
    if isinstance(value, str) and value.strip():
        for fmt in ("%Y-%m-%d", "%Y. %m. %d", "%Y.%m.%d", "%Y/%m/%d"):
            try:
                return datetime.strptime(value.strip()[:12].rstrip("."), fmt).date()
            except ValueError:
                pass
    return None


def rows_from_table(values: list[list]) -> list[dict]:
    if not values:
        return []
    header = [RENAMED.get(str(h), str(h)) for h in values[0]]
    rows = []
    for raw in values[1:]:
        row = {h: v for h, v in zip(header, raw) if h and h not in DROPPED}
        if row.get("ID"):
            row["ID"] = str(row["ID"])
            rows.append(row)
    return rows


def build_table(rows: list[dict], today: date | None = None) -> tuple[list[str], list[list]]:
    extra = []  # 사용자가 직접 추가한 열
    for r in rows:
        extra += [k for k in r if k not in HEADER and k not in extra]
    header = HEADER + extra
    far = date(9999, 12, 31)
    today = today or datetime.now(KST).date()

    def sort_key(r):  # 구분 → 진행 중(D-day 가까운 순) → 마감된 것은 맨 아래
        d = to_date(r.get("마감일")) or far
        return (str(r.get("구분", "")), d < today, d if d >= today else date.max - (d - date.min))

    rows = sorted(rows, key=sort_key)
    table = []
    for r in rows:
        line = []
        for h in header:
            v = r.get(h, "")
            if h == "D-day":
                v = DDAY_FORMULA
            elif h in DATE_COLS:
                v = to_date(v) or v
            line.append("" if v is None else v)
        table.append(line)
    return header, table


def merge_rows(existing: list[dict], new: list[dict]) -> list[dict]:
    """신규 행 추가 + 기존 행의 빈 칸(요약 등)만 보강. 사람이 적은 값은 덮어쓰지 않는다."""
    by_id = {r["ID"]: r for r in existing}
    for r in new:
        old = by_id.get(r["ID"])
        if old is None:
            by_id[r["ID"]] = r
        else:
            for k, v in r.items():
                if v not in (None, "") and old.get(k) in (None, ""):
                    old[k] = v
    return list(by_id.values())


class SheetStorage:
    def __init__(self, spreadsheet_id: str):
        import gspread

        raw = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON")
        if raw:
            gc = gspread.service_account_from_dict(json.loads(raw))
        else:
            gc = gspread.service_account(
                filename=os.environ.get("GOOGLE_SERVICE_ACCOUNT_FILE", "service_account.json"))
        self.sh = gc.open_by_key(spreadsheet_id)
        self.ws = self._get_or_create(MAIN_TAB, HEADER)
        self.seen_ws = self._get_or_create(SEEN_TAB, ["ID", "처리일", "결과"])
        self.settings_ws = self._settings_sheet()

    def _get_or_create(self, title: str, header: list[str], defaults: list[list] | None = None):
        import gspread

        try:
            ws = self.sh.worksheet(title)
        except gspread.WorksheetNotFound:
            ws = self.sh.add_worksheet(title=title, rows=1000, cols=max(len(header), 3))
        if not ws.row_values(1):
            ws.update([header] + (defaults or []), "A1")
            ws.freeze(rows=1)
        return ws

    def _settings_sheet(self):
        """'설정' 탭을 준비한다. 비어 있는 기본 시트(시트1)가 있으면 그걸 설정 탭으로 쓴다."""
        import gspread

        titles = [w.title for w in self.sh.worksheets()]
        blank = next((self.sh.worksheet(t) for t in ("시트1", "Sheet1")
                      if t in titles and not any(self.sh.worksheet(t).get_all_values())), None)
        if SETTINGS_TAB in titles:
            ws = self.sh.worksheet(SETTINGS_TAB)
            if blank is not None:
                self.sh.del_worksheet(blank)
        elif blank is not None:
            blank.update_title(SETTINGS_TAB)
            ws = blank
        else:
            ws = self.sh.add_worksheet(title=SETTINGS_TAB, rows=500, cols=8)
        try:
            ws.update_index(0)  # 맨 앞 탭으로
        except gspread.exceptions.APIError:
            pass
        return ws

    def seed_settings(self, config: dict) -> None:
        """처음 한 번(설명 표가 없을 때) config.yaml 의 모든 조건을 설정 탭으로 옮기고 설명 표를 붙인다."""
        values = self.settings_ws.get_all_values()
        if values and len(values[0]) >= 5 and values[0][4] == GUIDE_MARK:
            return
        order = list(SETTING_KINDS)
        rows = [r[:3] + [""] * (3 - len(r[:3])) for r in values[1:] if len(r) >= 2 and r[0].strip() and r[1].strip()]
        have = {(r[0].strip(), r[1].strip()) for r in rows}
        for kind, key in SETTING_KINDS.items():
            for v in (config.get(key, []) if key else []):
                if (kind, str(v)) not in have:
                    rows.append([kind, str(v), ""])
                    have.add((kind, str(v)))
        rows.sort(key=lambda r: order.index(r[0]) if r[0] in order else len(order))
        body = [SETTINGS_HEADER] + rows
        n = max(len(body), len(GUIDE))
        table = [(body[i] if i < len(body) else ["", "", ""]) + [""] +
                 (GUIDE[i] if i < len(GUIDE) else ["", "", ""]) for i in range(n)]
        self.settings_ws.resize(rows=max(self.settings_ws.row_count, n + 100), cols=max(self.settings_ws.col_count, 7))
        self.settings_ws.clear()
        self.settings_ws.update(table, "A1", value_input_option="RAW")
        self.settings_ws.freeze(rows=1)
        print(f"설정 탭 초기화: 조건 {len(rows)}개 + 설명 표")

    def load_rows(self) -> list[dict]:
        return rows_from_table(self.ws.get_all_values(value_render_option="UNFORMATTED_VALUE"))

    def seen_ids(self) -> set[str]:
        return {r["ID"] for r in self.load_rows()} | {str(i) for i in self.seen_ws.col_values(1)[1:] if i}

    def load_settings(self) -> dict[str, list[str]] | None:
        settings: dict[str, list[str]] = {}
        for row in self.settings_ws.get_all_values()[1:]:
            if len(row) >= 2 and row[0].strip() and row[1].strip():
                settings.setdefault(row[0].strip(), []).append(row[1].strip())
        return settings

    def save(self, rows: list[dict], seen: list[list]) -> None:
        header, table = build_table(rows)
        for line in table:
            for i, v in enumerate(line):
                if isinstance(v, date):
                    line[i] = v.isoformat()
        self.ws.resize(rows=max(self.ws.row_count, len(table) + 50), cols=max(self.ws.col_count, len(header)))
        self.ws.batch_clear(["A1:ZZ"])
        self.ws.update([header] + table, "A1", value_input_option="USER_ENTERED")
        if seen:
            self.seen_ws.append_rows(seen, value_input_option="RAW")
        self.apply_format(header)

    def apply_format(self, header: list[str]) -> None:
        """가독성 서식(sheet_format.py)을 다시 건다. 기존 줄무늬·조건부 서식은 지우고 새로 걸어 중복되지 않게 한다."""
        from sheet_format import build_requests

        try:
            meta = self.sh.fetch_sheet_metadata(params={
                "fields": "sheets.properties.sheetId,sheets.bandedRanges.bandedRangeId,sheets.conditionalFormats.ranges.sheetId"})
            info = next(x for x in meta["sheets"] if x["properties"]["sheetId"] == self.ws.id)
            banding = [b["bandedRangeId"] for b in info.get("bandedRanges", [])]
            n_cond = len(info.get("conditionalFormats", []))
            self.sh.batch_update({"requests": build_requests(self.ws.id, header, banding, n_cond)})
        except Exception as e:  # noqa: BLE001 - 서식 실패가 데이터 저장을 막지 않도록
            print(f"  ! 시트 서식 적용 실패(데이터는 저장됨): {type(e).__name__}: {e}")


class ExcelStorage:
    def __init__(self, path: Path):
        from openpyxl import Workbook, load_workbook

        self.path = path
        if path.exists():
            self.wb = load_workbook(path)
        else:
            path.parent.mkdir(parents=True, exist_ok=True)
            self.wb = Workbook()
            self.wb.active.title = MAIN_TAB
        self.ws = self._sheet(MAIN_TAB, HEADER)
        self.seen_ws = self._sheet(SEEN_TAB, ["ID", "처리일", "결과"])

    def _sheet(self, title: str, header: list[str]):
        ws = self.wb[title] if title in self.wb.sheetnames else self.wb.create_sheet(title)
        if ws.max_row == 1 and ws.cell(1, 1).value is None:
            for col, name in enumerate(header, start=1):
                ws.cell(1, col, name)
            ws.freeze_panes = "A2"
        return ws

    def load_rows(self) -> list[dict]:
        return rows_from_table([list(r) for r in self.ws.iter_rows(values_only=True)])

    def seen_ids(self) -> set[str]:
        ids = {r["ID"] for r in self.load_rows()}
        for (value,) in self.seen_ws.iter_rows(min_row=2, max_col=1, values_only=True):
            if value:
                ids.add(str(value))
        return ids

    def load_settings(self) -> dict[str, list[str]] | None:
        return None  # 엑셀 모드는 config.yaml 만 사용

    def save(self, rows: list[dict], seen: list[list]) -> None:
        header, table = build_table(rows)
        self.ws.delete_rows(1, self.ws.max_row)
        self.ws.append(header)
        link_col = header.index("링크") + 1
        for i, line in enumerate(table, start=2):
            self.ws.append(line)
            for name in DATE_COLS:
                self.ws.cell(i, header.index(name) + 1).number_format = "yyyy-mm-dd"
            if line[link_col - 1]:
                self.ws.cell(i, link_col).hyperlink = line[link_col - 1]
        for s in seen:
            self.seen_ws.append(s)
        self.wb.save(self.path)
